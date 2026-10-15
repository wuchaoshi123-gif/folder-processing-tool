# export_failed_xlsx.py
"""导出失败记录为 xlsx（含图片路径超链接）"""
import sqlite3
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


DB_PATH = r"L:\PythonProject1\Folder_processing_tool\data\products.db"


def main():
    p = Path(DB_PATH)
    if not p.exists():
        print(f"❌ 数据库不存在: {p}")
        return

    conn = sqlite3.connect(str(p))
    cur = conn.cursor()

    cur.execute("""
        SELECT id, onlymulu, reason, detail, image_path, ocr_text, created_at
        FROM failed_items
        ORDER BY reason, onlymulu
    """)
    rows = cur.fetchall()
    conn.close()

    if not rows:
        print("✅ 无失败记录")
        return

    print(f"📊 共 {len(rows)} 条失败记录")

    # 写 xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = "失败记录"

    headers = ["ID", "唯一目录", "失败原因", "详细原因",
               "图片路径", "OCR 原始文本", "时间"]
    ws.append(headers)

    # 表头样式
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.alignment = Alignment(horizontal="center")

    # 数据行
    for r in rows:
        rid, om, reason, detail, img_path, ocr_text, ca = r
        ws.append([
            rid, om or "", reason or "", detail or "",
            img_path or "", ocr_text or "", ca or ""
        ])
        # ★ 图片路径加超链接
        if img_path:
            img_cell = ws.cell(row=ws.max_row, column=5)
            try:
                img_cell.hyperlink = f"file:///{img_path.replace(chr(92), '/')}"
                img_cell.font = Font(color="0563C1", underline="single")
            except Exception:
                pass

    # 列宽
    widths = [6, 50, 20, 40, 70, 60, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 冻结首行
    ws.freeze_panes = "A2"

    # 保存
    out = p.parent / f"failed_items_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(out)
    print(f"✅ 已导出: {out}")
    print(f"\n💡 提示: 用 Excel 打开后，点击【图片路径】列可直接打开图片")


if __name__ == "__main__":
    main()