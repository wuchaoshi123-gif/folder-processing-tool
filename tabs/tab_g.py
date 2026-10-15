# tabs/tab_g.py
"""G 菜单：产品表格生成（v6 使用标记版）
★ v6:
  - 新增 g_used_products 表，随机抽取自动排除已用
  - 显示剩余未用数量
  - 生成成功后自动标记（可开关）
  - 一键清空标记
"""
import json
import os
import threading
import tkinter as tk
from pathlib import Path
from datetime import datetime
from tkinter import ttk, messagebox, filedialog
import sqlite3

from .base_tab import BaseTab
from core import g_generator as gg


APP_NAME = "FolderProcessingTool"
AUTO_CFG_DIR = Path(os.getenv("APPDATA", Path.home())) / APP_NAME
AUTO_CFG_DIR.mkdir(parents=True, exist_ok=True)
AUTO_CFG_PATH = AUTO_CFG_DIR / "tab_g_config.json"

LEGACY_PRICE_EXPRS = {
    "CHUPRICE*2+rand(50,90)",
    "CHUPRICE*2+rand(99,160)",
    "CHUPRICE*2+rand(150,199)",
    "CHUPRICE*2+rand(160,199)",
    "CHUPRICE*2.10+rand(66,99)",
    "round((CHUPRICE * 2.10 + rand(66, 99)) / (0.75 * (1 - 0.17)))",
}


class TabG(BaseTab):
    TAB_NAME = "G 产品表格"
    TAB_KEY = "G"

    # ==========================================================
    #                    UI 构建
    # ==========================================================
    def _build_ui(self):
        self._conn = None
        self._field_rows = {"A": [], "B": [], "C": []}
        self._field_container = {}
        self._loading = False
        self._save_job = None

        # ==========================================================
        # ★ 2×2 田字格容器
        #   上左: 数据源            上右: 模板 & 输出
        #   下左: 参数              下右: 价格公式参数
        # ==========================================================
        grid_frame = ttk.Frame(self)
        grid_frame.pack(fill="x", padx=6, pady=(6, 4))
        grid_frame.columnconfigure(0, weight=1, uniform="g_col")
        grid_frame.columnconfigure(1, weight=1, uniform="g_col")

        # ==========================================================
        # 上左：数据源
        # ==========================================================
        src = ttk.LabelFrame(grid_frame, text=" 数据源 ", padding=6)
        src.grid(row=0, column=0, sticky="nsew", padx=(0, 4), pady=(0, 4))

        row1 = ttk.Frame(src)
        row1.pack(fill="x")
        ttk.Label(row1, text="数据库:").pack(side="left")
        self.db_var = tk.StringVar()
        ttk.Entry(row1, textvariable=self.db_var).pack(
            side="left", fill="x", expand=True, padx=4)
        ttk.Button(row1, text="浏览", width=5,
                   command=self._choose_db).pack(side="left", padx=2)
        ttk.Button(row1, text="刷新", width=5,
                   command=self._reload_db).pack(side="left", padx=2)

        row1b = ttk.Frame(src)
        row1b.pack(fill="x", pady=(2, 0))
        self.db_info_var = tk.StringVar(value="未连接")
        ttk.Label(row1b, textvariable=self.db_info_var,
                  foreground="#004080").pack(side="left")

        row2 = ttk.Frame(src)
        row2.pack(fill="x", pady=(6, 0))
        ttk.Label(row2, text="抽取:").pack(side="left")
        self.db_sample_mode_var = tk.StringVar(value="全部")
        ttk.Radiobutton(row2, text="全部",
                        variable=self.db_sample_mode_var,
                        value="全部",
                        command=self._on_param_change).pack(side="left", padx=4)
        ttk.Radiobutton(row2, text="随机",
                        variable=self.db_sample_mode_var,
                        value="随机抽取",
                        command=self._on_param_change).pack(side="left", padx=4)
        ttk.Label(row2, text="数量:").pack(side="left", padx=(8, 0))
        self.db_sample_count_var = tk.StringVar(value="100")
        ttk.Spinbox(row2, from_=1, to=1000000, increment=10,
                    textvariable=self.db_sample_count_var,
                    width=7,
                    command=self._on_param_change).pack(side="left", padx=4)

        row2b = ttk.Frame(src)
        row2b.pack(fill="x", pady=(2, 0))
        self.db_sample_info_var = tk.StringVar(value="")
        ttk.Label(row2b, textvariable=self.db_sample_info_var,
                  foreground="#804000",
                  font=("Consolas", 9, "bold")).pack(side="left")

        row3 = ttk.Frame(src)
        row3.pack(fill="x", pady=(6, 0))
        self.mark_used_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(row3,
                        text="生成成功后自动标记",
                        variable=self.mark_used_var).pack(side="left")
        self.used_info_var = tk.StringVar(value="")
        ttk.Label(row3, textvariable=self.used_info_var,
                  foreground="#666",
                  font=("Consolas", 8)).pack(side="left", padx=4)
        ttk.Button(row3, text="🗑 清空标记", width=11,
                   command=self._clear_used_marks).pack(side="right", padx=2)

        # ==========================================================
        # 上右：模板 & 输出
        # ==========================================================
        tmp = ttk.LabelFrame(grid_frame, text=" 模板 & 输出 ", padding=6)
        tmp.grid(row=0, column=1, sticky="nsew", padx=(4, 0), pady=(0, 4))

        ttk.Label(tmp, text="模板:").grid(row=0, column=0, sticky="w", pady=2)
        self.tpl_var = tk.StringVar(value="休闲运动鞋.xlsx")
        ttk.Entry(tmp, textvariable=self.tpl_var).grid(
            row=0, column=1, sticky="ew", padx=4)
        ttk.Button(tmp, text="浏览", width=6,
                   command=self._choose_template).grid(
            row=0, column=2, padx=2)

        ttk.Label(tmp, text="输出目录:").grid(row=1, column=0, sticky="w", pady=2)
        self.out_var = tk.StringVar()
        ttk.Entry(tmp, textvariable=self.out_var).grid(
            row=1, column=1, sticky="ew", padx=4)
        ttk.Button(tmp, text="浏览", width=6,
                   command=self._choose_output).grid(
            row=1, column=2, padx=2)

        ttk.Label(tmp, text="关键词:").grid(row=2, column=0, sticky="w", pady=2)
        self.kw_var = tk.StringVar(value="关键词.txt")
        ttk.Entry(tmp, textvariable=self.kw_var).grid(
            row=2, column=1, sticky="ew", padx=4)
        ttk.Button(tmp, text="浏览", width=6,
                   command=self._choose_keyword).grid(
            row=2, column=2, padx=2)

        tmp.columnconfigure(1, weight=1)

        # ==========================================================
        # 下左：参数
        # ==========================================================
        par = ttk.LabelFrame(grid_frame, text=" 参数 ", padding=6)
        par.grid(row=1, column=0, sticky="nsew", padx=(0, 4), pady=(4, 0))

        # 第 1 行：列名行 / 数据起始行 / 单文件最多
        ttk.Label(par, text="列名行:").grid(row=0, column=0, sticky="w", pady=2)
        self.header_row_var = tk.StringVar(value="3")
        ttk.Spinbox(par, from_=1, to=20, width=4,
                    textvariable=self.header_row_var,
                    command=self._on_param_change).grid(
            row=0, column=1, sticky="w", padx=2)

        ttk.Label(par, text="数据起始:").grid(
            row=0, column=2, sticky="w", padx=(10, 0))
        self.start_row_var = tk.StringVar(value="5")
        ttk.Spinbox(par, from_=1, to=20, width=4,
                    textvariable=self.start_row_var,
                    command=self._on_param_change).grid(
            row=0, column=3, sticky="w", padx=2)

        ttk.Label(par, text="单文件最多:").grid(
            row=0, column=4, sticky="w", padx=(10, 0))
        self.max_prod_var = tk.StringVar(value="950")
        ttk.Spinbox(par, from_=1, to=1000, width=5,
                    textvariable=self.max_prod_var,
                    command=self._on_param_change).grid(
            row=0, column=5, sticky="w", padx=2)

        # 第 2 行：图片后缀 + 首图
        ttk.Label(par, text="图片后缀:").grid(row=1, column=0, sticky="w", pady=2)
        self.img_suffix_var = tk.StringVar(
            value="?x-oss-process=style/fnew06")
        ttk.Entry(par, textvariable=self.img_suffix_var).grid(
            row=1, column=1, columnspan=4, sticky="ew", padx=2)

        ttk.Label(par, text="首图:").grid(
            row=1, column=5, sticky="w", padx=(4, 0))
        self.first_img_var = tk.StringVar(value="1.jpg")
        ttk.Entry(par, textvariable=self.first_img_var, width=8).grid(
            row=1, column=6, sticky="w", padx=2)

        # 第 3 行：抽图数量 / 名称最大长度
        ttk.Label(par, text="抽图数量:").grid(row=2, column=0, sticky="w", pady=2)
        self.sample_count_var = tk.StringVar(value="0")
        ttk.Spinbox(par, from_=0, to=100, width=4,
                    textvariable=self.sample_count_var,
                    command=self._on_param_change).grid(
            row=2, column=1, sticky="w", padx=2)
        ttk.Label(par, text="(0=全部)",
                  foreground="#666").grid(row=2, column=2, sticky="w")

        ttk.Label(par, text="名称长度:").grid(
            row=2, column=3, sticky="w", padx=(4, 0))
        self.name_len_var = tk.StringVar(value="60")
        ttk.Spinbox(par, from_=10, to=200, width=5,
                    textvariable=self.name_len_var,
                    command=self._on_param_change).grid(
            row=2, column=4, sticky="w", padx=2)

        # 第 4 行：关键词位置
        ttk.Label(par, text="关键词位置:").grid(
            row=3, column=0, sticky="w", pady=2)
        self.kw_pos_var = tk.StringVar(value="random")
        ttk.Combobox(par, textvariable=self.kw_pos_var, width=10,
                     state="readonly",
                     values=["random", "before", "after"]).grid(
            row=3, column=1, sticky="w", padx=2)

        par.columnconfigure(1, weight=1)

        # ==========================================================
        # 下右：价格公式参数
        # ==========================================================
        price = ttk.LabelFrame(grid_frame, text=" 💰 价格公式参数 ", padding=6)
        price.grid(row=1, column=1, sticky="nsew", padx=(4, 0), pady=(4, 0))

        # 目标利润
        ttk.Label(price, text="目标利润:").grid(
            row=0, column=0, sticky="w", pady=2)
        ttk.Label(price, text="成本 ×", foreground="#666").grid(
            row=0, column=1, sticky="w")
        self.profit_rate_var = tk.StringVar(value="110")
        ttk.Spinbox(price, from_=0, to=500, increment=5, width=5,
                    textvariable=self.profit_rate_var,
                    command=self._on_param_change).grid(
            row=0, column=2, sticky="w", padx=2)
        ttk.Label(price, text="%").grid(row=0, column=3, sticky="w")

        # 售价折扣
        ttk.Label(price, text="售价折扣:").grid(
            row=1, column=0, sticky="w", pady=2)
        self.discount_var = tk.StringVar(value="0.75")
        ttk.Entry(price, textvariable=self.discount_var, width=7).grid(
            row=1, column=1, sticky="w", padx=2)
        ttk.Label(price, text="(如 0.75=打75折)",
                  foreground="#666", font=("", 8)).grid(
            row=1, column=2, columnspan=2, sticky="w")

        # 平台费率
        ttk.Label(price, text="平台费率:").grid(
            row=2, column=0, sticky="w", pady=2)
        self.platform_fee_var = tk.StringVar(value="0.17")
        ttk.Entry(price, textvariable=self.platform_fee_var, width=7).grid(
            row=2, column=1, sticky="w", padx=2)
        ttk.Label(price, text="(如 0.17=17%)",
                  foreground="#666", font=("", 8)).grid(
            row=2, column=2, columnspan=2, sticky="w")

        # 运费范围
        ttk.Label(price, text="运费范围:").grid(
            row=3, column=0, sticky="w", pady=2)
        self.freight_min_var = tk.StringVar(value="66")
        ttk.Spinbox(price, from_=0, to=10000, width=5,
                    textvariable=self.freight_min_var,
                    command=self._on_param_change).grid(
            row=3, column=1, sticky="w", padx=2)
        ttk.Label(price, text="~").grid(row=3, column=2, sticky="w")
        self.freight_max_var = tk.StringVar(value="99")
        ttk.Spinbox(price, from_=0, to=10000, width=5,
                    textvariable=self.freight_max_var,
                    command=self._on_param_change).grid(
            row=3, column=3, sticky="w", padx=2)

        # 示例成本
        ttk.Label(price, text="示例成本:").grid(
            row=4, column=0, sticky="w", pady=2)
        self.sample_cost_var = tk.StringVar(value="150")
        ttk.Spinbox(price, from_=0, to=100000, width=7,
                    textvariable=self.sample_cost_var,
                    command=self._on_param_change).grid(
            row=4, column=1, sticky="w", padx=2)
        ttk.Label(price, text="（仅影响下方示例）",
                  foreground="#666", font=("", 8)).grid(
            row=4, column=2, columnspan=2, sticky="w")

        # 应用按钮
        ttk.Button(price, text="🔧 应用到 A 组【价格】", width=22,
                   command=self._apply_price_formula).grid(
            row=5, column=0, columnspan=4, sticky="w", pady=(6, 0))

        # 示例显示
        self.price_preview_var = tk.StringVar(value="")
        ttk.Label(price, textvariable=self.price_preview_var,
                  foreground="#804000",
                  font=("Consolas", 9, "bold")).grid(
            row=6, column=0, columnspan=4, sticky="w", pady=(2, 0))

        # 价格参数 trace
        for v in (self.profit_rate_var, self.discount_var,
                  self.platform_fee_var, self.freight_min_var,
                  self.freight_max_var, self.sample_cost_var):
            try:
                v.trace_add("write", lambda *a: (
                    self._update_price_preview(), self._schedule_save()))
            except Exception:
                pass

        # ==========================================================
        # 字段编辑器（占满剩余高度）
        # ==========================================================
        fields_frame = ttk.LabelFrame(self, text=" 字段定义 ", padding=6)
        fields_frame.pack(fill="both", expand=True, padx=6, pady=(0, 4))

        nb = ttk.Notebook(fields_frame)
        nb.pack(fill="both", expand=True)

        for kind, title, hint in [
            ("A", "A 组（每组一次）", "产品级字段，组内所有行共用"),
            ("B", "B 组（每行一次）", "行级字段，SIZE / SIZE_INT 可用"),
            ("C", "C 组（固定值）", "全表固定，直接写死值"),
        ]:
            fr = ttk.Frame(nb, padding=6)
            nb.add(fr, text=title)
            ttk.Label(fr, text=hint, foreground="#666").pack(
                anchor="w", pady=(0, 4))

            body = ttk.Frame(fr)
            body.pack(fill="both", expand=True)

            head = ttk.Frame(body)
            head.pack(fill="x")
            ttk.Label(head, text="xlsx 列名", width=20,
                      font=("", 9, "bold")).pack(side="left")
            ttk.Label(head, text="表达式 / 值",
                      font=("", 9, "bold")).pack(side="left", padx=6)

            listbox = ttk.Frame(body)
            listbox.pack(fill="both", expand=True, pady=4)
            self._field_container[kind] = listbox

            btn = ttk.Frame(fr)
            btn.pack(fill="x", pady=(4, 0))
            ttk.Button(btn, text=f"+ 新增 {kind} 字段",
                       command=lambda k=kind: self._add_field(k)
                       ).pack(side="left")

        # ==========================================================
        # 操作按钮
        # ==========================================================
        op = ttk.Frame(self, padding=(6, 0, 6, 4))
        op.pack(fill="x")

        ttk.Button(op, text="📥 加载模板列名", width=16,
                   command=self._load_template_columns).pack(side="left", padx=2)
        ttk.Button(op, text="👁 预览前 3 组", width=16,
                   command=self._preview).pack(side="left", padx=2)
        ttk.Button(op, text="▶ 生成 xlsx", width=16,
                   command=self._generate).pack(side="left", padx=2)
        ttk.Button(op, text="✖ 清空", width=8,
                   command=self._clear).pack(side="left", padx=2)

        ttk.Separator(op, orient="vertical").pack(side="left", fill="y", padx=6)

        ttk.Button(op, text="💾 导出配置", width=14,
                   command=self._export_config).pack(side="left", padx=2)
        ttk.Button(op, text="📂 导入配置", width=14,
                   command=self._import_config).pack(side="left", padx=2)

        self.msg_var = tk.StringVar(value="")
        ttk.Label(op, textvariable=self.msg_var,
                  foreground="#804000").pack(side="left", padx=10)

        # ==========================================================
        # 预览文本区
        # ==========================================================
        self.preview_text = tk.Text(self, height=10, font=("Consolas", 9),
                                    wrap="none")
        self.preview_text.pack(fill="both", expand=True, padx=6, pady=(0, 6))
        sb = ttk.Scrollbar(self.preview_text, command=self.preview_text.yview)
        sb.pack(side="right", fill="y")
        self.preview_text.config(yscrollcommand=sb.set)

        # 延迟初始化
        self.app.root.after(300, self._post_build)

    # ==========================================================
    def _post_build(self):
        self._load_auto_config()
        self._bind_traces()
        self._update_price_preview()
        try:
            self.app.root.protocol("WM_DELETE_WINDOW", self._on_close)
        except Exception:
            pass
        self.app.root.after(800, self._auto_load_db)

    def _bind_traces(self):
        for var in [
            self.db_var, self.tpl_var, self.out_var, self.kw_var,
            self.header_row_var, self.start_row_var, self.max_prod_var,
            self.img_suffix_var, self.first_img_var, self.sample_count_var,
            self.name_len_var, self.kw_pos_var,
            self.db_sample_mode_var, self.db_sample_count_var,
        ]:
            try:
                var.trace_add("write", lambda *a: self._schedule_save())
            except Exception:
                pass

    def _on_param_change(self):
        self._update_sample_info()
        self._update_price_preview()
        self._schedule_save()

    def _on_close(self):
        try:
            self._save_auto_config()
        except Exception:
            pass
        self.app.root.destroy()

    # ==========================================================
    #  ★ 已用标记管理
    # ==========================================================
    def _ensure_used_table(self):
        """确保 g_used_products 表存在"""
        if not self._conn:
            return
        try:
            cur = self._conn.cursor()
            cur.execute("""
                CREATE TABLE IF NOT EXISTS g_used_products (
                    onlymulu TEXT PRIMARY KEY,
                    used_at  DATETIME DEFAULT CURRENT_TIMESTAMP
                )
            """)
            self._conn.commit()
        except Exception as e:
            print(f"[G] 建表失败: {e}")

    def _used_count(self):
        if not self._conn:
            return 0
        try:
            cur = self._conn.cursor()
            cur.execute("SELECT COUNT(*) FROM g_used_products")
            return cur.fetchone()[0]
        except Exception:
            return 0

    def _mark_used(self, onlymulu_list):
        """批量标记为已用，返回新标记数"""
        if not onlymulu_list or not self._conn:
            return 0
        try:
            cur = self._conn.cursor()
            cnt = 0
            # 分批，避免参数过多
            for i in range(0, len(onlymulu_list), 500):
                batch = onlymulu_list[i:i + 500]
                ph = ",".join(["(?)"] * len(batch))
                cur.execute(
                    f"INSERT OR IGNORE INTO g_used_products (onlymulu) "
                    f"VALUES {ph}", batch)
                cnt += cur.rowcount
            self._conn.commit()
            return cnt
        except Exception as e:
            print(f"[G] 标记失败: {e}")
            return 0

    def _clear_used_marks(self):
        if not self._conn:
            messagebox.showinfo("提示", "数据库未连接")
            return
        used = self._used_count()
        if used == 0:
            messagebox.showinfo("提示", "没有需要清空的标记")
            return
        if not messagebox.askyesno(
            "确认清空",
            f"将清除 {used:,} 条已用标记\n\n"
            f"清空后所有产品都能再次被随机抽取。\n\n继续?"):
            return
        try:
            cur = self._conn.cursor()
            cur.execute("DELETE FROM g_used_products")
            self._conn.commit()
            self._update_sample_info()
            messagebox.showinfo("完成", f"已清除 {used:,} 条标记")
        except Exception as e:
            messagebox.showerror("清空失败", str(e))

    # ==========================================================
    def _update_price_preview(self):
        if not hasattr(self, "price_preview_var"):
            return
        try:
            pr = float(self.profit_rate_var.get() or 110) / 100.0
            dc = float(self.discount_var.get() or 0.75)
            pf = float(self.platform_fee_var.get() or 0.17)
            fmin = int(float(self.freight_min_var.get() or 66))
            fmax = int(float(self.freight_max_var.get() or 99))
            sample_cost = float(self.sample_cost_var.get() or 100)
        except (ValueError, AttributeError):
            self.price_preview_var.set("⚠️ 参数格式错误")
            return

        if dc <= 0 or dc > 1 or pf < 0 or pf >= 1:
            self.price_preview_var.set("⚠️ 折扣应 0~1；费率应 0~1")
            return

        denom = dc * (1 - pf)
        if denom <= 0:
            self.price_preview_var.set("⚠️ 折扣或费率为 0，公式无效")
            return

        freight_mid = (fmin + fmax) // 2
        price_mid = round((sample_cost * (1 + pr) + freight_mid) / denom)
        price_lo = round((sample_cost * (1 + pr) + fmin) / denom)
        price_hi = round((sample_cost * (1 + pr) + fmax) / denom)

        self.price_preview_var.set(
            f"→ 示例: 成本{int(sample_cost)}, 运费{fmin}~{fmax} 时, 售价 ≈ "
            f"{price_lo} ~ {price_hi} (中值 {price_mid})")

    def _apply_price_formula(self):
        target_col = "价格"
        new_expr = "calc_price(CHUPRICE)"
        found = False
        for col_var, expr_var, _ in self._field_rows["A"]:
            if col_var.get().strip() == target_col:
                expr_var.set(new_expr)
                found = True
                break
        if not found:
            self._add_field("A", target_col, new_expr)
        self._schedule_save()
        messagebox.showinfo(
            "已应用",
            f"A 组【{target_col}】字段的表达式已设为:\n\n"
            f"    {new_expr}\n\n"
            f"现在点击【👁 预览前 3 组】看效果。")

    # ==========================================================
    #  字段行管理
    # ==========================================================
    def _add_field(self, kind, col="", expr="", skip_save=False):
        container = self._field_container[kind]
        row = ttk.Frame(container)
        row.pack(fill="x", pady=1)

        col_var = tk.StringVar(value=col)
        expr_var = tk.StringVar(value=expr)
        ttk.Entry(row, textvariable=col_var, width=20).pack(side="left")
        ttk.Entry(row, textvariable=expr_var).pack(
            side="left", fill="x", expand=True, padx=6)

        def remove():
            for i, item in enumerate(self._field_rows[kind]):
                if item[2] is row:
                    self._field_rows[kind].pop(i)
                    break
            row.destroy()
            self._schedule_save()

        ttk.Button(row, text="✖", width=3, command=remove).pack(side="left")

        try:
            col_var.trace_add("write", lambda *a: self._schedule_save())
            expr_var.trace_add("write", lambda *a: self._schedule_save())
        except Exception:
            pass

        self._field_rows[kind].append((col_var, expr_var, row))
        if not skip_save:
            self._schedule_save()

    def _clear_fields(self, kind):
        for _, _, row in self._field_rows[kind]:
            row.destroy()
        self._field_rows[kind] = []

    def _collect_fields(self, kind):
        return [
            (cv.get().strip(), ev.get().strip())
            for cv, ev, _ in self._field_rows[kind]
            if cv.get().strip()
        ]

    # ==========================================================
    #  配置收集 / 应用
    # ==========================================================
    def _collect_full_config(self):
        return {
            "db_path": self.db_var.get().strip(),
            "template_path": self.tpl_var.get().strip(),
            "output_dir": self.out_var.get().strip(),
            "keyword_file": self.kw_var.get().strip(),
            "header_row": self.header_row_var.get(),
            "data_start_row": self.start_row_var.get(),
            "max_products_per_file": self.max_prod_var.get(),
            "img_suffix": self.img_suffix_var.get(),
            "first_img": self.first_img_var.get(),
            "sample_count": self.sample_count_var.get(),
            "name_max_len": self.name_len_var.get(),
            "keyword_pos": self.kw_pos_var.get(),
            "db_sample_mode": self.db_sample_mode_var.get(),
            "db_sample_count": self.db_sample_count_var.get(),
            "mark_used": self.mark_used_var.get(),
            "profit_rate": self.profit_rate_var.get(),
            "discount": self.discount_var.get(),
            "platform_fee": self.platform_fee_var.get(),
            "freight_min": self.freight_min_var.get(),
            "freight_max": self.freight_max_var.get(),
            "sample_cost": self.sample_cost_var.get(),
            "fields_a": self._collect_fields("A"),
            "fields_b": self._collect_fields("B"),
            "fields_c": self._collect_fields("C"),
        }

    def _apply_full_config(self, cfg):
        if not cfg:
            return
        self._loading = True
        try:
            self.db_var.set(cfg.get("db_path", ""))
            self.tpl_var.set(cfg.get("template_path", "平板鞋.xlsx"))
            self.out_var.set(cfg.get("output_dir", ""))
            self.kw_var.set(cfg.get("keyword_file", "关键词.txt"))
            self.header_row_var.set(str(cfg.get("header_row", 3)))
            self.start_row_var.set(str(cfg.get("data_start_row", 5)))
            self.max_prod_var.set(str(cfg.get("max_products_per_file", 100)))
            self.img_suffix_var.set(cfg.get("img_suffix", ""))
            self.first_img_var.set(cfg.get("first_img", "1.jpg"))
            self.sample_count_var.set(str(cfg.get("sample_count", 0)))
            self.name_len_var.set(str(cfg.get("name_max_len", 60)))
            self.kw_pos_var.set(cfg.get("keyword_pos", "random"))
            self.db_sample_mode_var.set(cfg.get("db_sample_mode", "全部"))
            self.db_sample_count_var.set(str(cfg.get("db_sample_count", 100)))
            self.mark_used_var.set(bool(cfg.get("mark_used", True)))

            pr = cfg.get("profit_rate", 110)
            try:
                prf = float(pr)
                if prf <= 5:
                    prf = prf * 100
                self.profit_rate_var.set(str(int(prf)))
            except (ValueError, TypeError):
                self.profit_rate_var.set("110")
            self.discount_var.set(str(cfg.get("discount", 0.75)))
            self.platform_fee_var.set(str(cfg.get("platform_fee", 0.17)))
            self.freight_min_var.set(str(cfg.get("freight_min", 66)))
            self.freight_max_var.set(str(cfg.get("freight_max", 99)))
            self.sample_cost_var.set(str(cfg.get("sample_cost", 100)))

            for kind in ("A", "B", "C"):
                self._clear_fields(kind)
                for col, expr in cfg.get(f"fields_{kind.lower()}", []):
                    if kind == "A" and col == "价格" \
                            and expr in LEGACY_PRICE_EXPRS:
                        expr = "calc_price(CHUPRICE)"
                    self._add_field(kind, col, expr, skip_save=True)
        finally:
            self._loading = False
            self._update_sample_info()
            self._update_price_preview()

    # ==========================================================
    #  自动保存 / 加载
    # ==========================================================
    def _schedule_save(self):
        if self._loading:
            return
        if self._save_job:
            try:
                self.app.root.after_cancel(self._save_job)
            except Exception:
                pass
        self._save_job = self.app.root.after(800, self._save_auto_config)

    def _save_auto_config(self):
        self._save_job = None
        try:
            cfg = self._collect_full_config()
            AUTO_CFG_PATH.write_text(
                json.dumps(cfg, ensure_ascii=False, indent=2),
                encoding="utf-8")
        except Exception as e:
            print(f"[G] 自动保存失败: {e}")

    def _load_auto_config(self):
        if AUTO_CFG_PATH.exists():
            try:
                cfg = json.loads(AUTO_CFG_PATH.read_text(encoding="utf-8"))
                self._apply_full_config(cfg)
                print(f"[G] ✅ 已加载配置: {AUTO_CFG_PATH}")
                return
            except Exception as e:
                print(f"[G] ⚠️ 加载配置失败: {e}")

        for kind, key in [("A", "fields_a"), ("B", "fields_b"),
                          ("C", "fields_c")]:
            for col, expr in gg.DEFAULT_G_CONFIG[key]:
                self._add_field(kind, col, expr, skip_save=True)
        self.out_var.set(str(Path.cwd() / "exports"))

    # ==========================================================
    #  导入 / 导出
    # ==========================================================
    def _export_config(self):
        p = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON", "*.json"), ("All", "*.*")],
            initialfile=f"tab_g_config_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json",
            title="导出 G 模块配置")
        if not p:
            return
        try:
            cfg = self._collect_full_config()
            Path(p).write_text(
                json.dumps(cfg, ensure_ascii=False, indent=2),
                encoding="utf-8")
            messagebox.showinfo("导出成功", f"已保存到:\n{p}")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _import_config(self):
        p = filedialog.askopenfilename(
            filetypes=[("JSON", "*.json"), ("All", "*.*")],
            title="选择要导入的 G 模块配置")
        if not p:
            return
        try:
            cfg = json.loads(Path(p).read_text(encoding="utf-8"))
            self._apply_full_config(cfg)
            self._save_auto_config()
            messagebox.showinfo("导入成功", "配置已应用并保存")
        except Exception as e:
            messagebox.showerror("导入失败", str(e))

    # ==========================================================
    #  数据库
    # ==========================================================
    def _auto_load_db(self):
        try:
            from config_manager import load_config
            cfg = load_config()
            db = cfg.get("DB_PATH", "")
            if db and Path(db).exists() and not self.db_var.get().strip():
                self.db_var.set(db)
            if self.db_var.get().strip():
                self._reload_db()
        except Exception as e:
            print(f"[G] 自动加载失败: {e}")

    def _choose_db(self):
        p = filedialog.askopenfilename(
            filetypes=[("SQLite DB", "*.db"), ("All", "*.*")])
        if p:
            self.db_var.set(p)
            self._reload_db()

    def _choose_template(self):
        p = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx"), ("All", "*.*")])
        if p:
            self.tpl_var.set(p)

    def _choose_output(self):
        p = filedialog.askdirectory()
        if p:
            self.out_var.set(p)

    def _choose_keyword(self):
        p = filedialog.askopenfilename(
            filetypes=[("Text", "*.txt"), ("All", "*.*")])
        if p:
            self.kw_var.set(p)

    def _reload_db(self):
        db = self.db_var.get().strip()
        if not db or not Path(db).exists():
            self.db_info_var.set("数据库不存在")
            return
        try:
            if self._conn:
                self._conn.close()
            self._conn = sqlite3.connect(db)
            self._conn.row_factory = sqlite3.Row
            self._ensure_used_table()       # ★ 确保标记表存在
            n = self._total_products()
            self.db_info_var.set(f"✅ {n:,} 条")
            self._update_sample_info()
        except Exception as e:
            self.db_info_var.set(f"❌ {str(e)[:40]}")

    def _total_products(self):
        try:
            cur = self._conn.cursor()
            cur.execute("SELECT COUNT(*) FROM products")
            return cur.fetchone()[0]
        except Exception:
            return 0

    def _get_products(self, limit=None):
        if not self._conn:
            self._reload_db()
        if not self._conn:
            return []
        cur = self._conn.cursor()

        if self.db_sample_mode_var.get() == "随机抽取":
            try:
                n = int(self.db_sample_count_var.get() or 100)
            except ValueError:
                n = 100
            if n <= 0:
                n = 100

            # ★ 从"未用"里抽
            cur.execute(
                "SELECT * FROM products "
                "WHERE onlymulu NOT IN "
                "  (SELECT onlymulu FROM g_used_products) "
                "ORDER BY RANDOM() LIMIT ?", (n,))
        else:
            cur.execute("SELECT * FROM products ORDER BY id")

        rows = [dict(r) for r in cur.fetchall()]
        if limit:
            rows = rows[:limit]
        return rows

    def _update_sample_info(self):
        if not self._conn:
            return
        total = self._total_products()
        used = self._used_count()
        remaining = total - used

        mode = self.db_sample_mode_var.get()
        if mode == "随机抽取":
            try:
                n = int(self.db_sample_count_var.get() or 100)
            except ValueError:
                n = 100
            actual = min(n, remaining)
            self.db_sample_info_var.set(
                f"→ 本次抽 {actual:,} / 剩余未用 {remaining:,} / 共 {total:,}")
        else:
            self.db_sample_info_var.set(
                f"→ 全部 {total:,} 条")

        if used > 0:
            self.used_info_var.set(f"（已用 {used:,}）")
        else:
            self.used_info_var.set("")

    def _build_config(self):
        try:
            pr = float(self.profit_rate_var.get() or 110) / 100.0
        except ValueError:
            pr = 1.10
        try:
            dc = float(self.discount_var.get() or 0.75)
        except ValueError:
            dc = 0.75
        try:
            pf = float(self.platform_fee_var.get() or 0.17)
        except ValueError:
            pf = 0.17
        try:
            fmin = int(float(self.freight_min_var.get() or 66))
        except ValueError:
            fmin = 66
        try:
            fmax = int(float(self.freight_max_var.get() or 99))
        except ValueError:
            fmax = 99

        return {
            "template_path": self.tpl_var.get().strip(),
            "output_dir": self.out_var.get().strip(),
            "header_row": int(self.header_row_var.get() or 3),
            "data_start_row": int(self.start_row_var.get() or 5),
            "max_products_per_file": int(self.max_prod_var.get() or 100),
            "img_suffix": self.img_suffix_var.get(),
            "first_img": self.first_img_var.get().strip(),
            "sample_count": int(self.sample_count_var.get() or 0),
            "img_mode": "single",
            "keyword_file": self.kw_var.get().strip(),
            "keyword_pos": self.kw_pos_var.get().strip(),
            "name_max_len": int(self.name_len_var.get() or 60),
            "profit_rate": pr,
            "discount": dc,
            "platform_fee": pf,
            "freight_min": fmin,
            "freight_max": fmax,
            "fields_a": self._collect_fields("A"),
            "fields_b": self._collect_fields("B"),
            "fields_c": self._collect_fields("C"),
        }

    # ==========================================================
    def _load_template_columns(self):
        tpl = self.tpl_var.get().strip()
        if not Path(tpl).exists():
            messagebox.showerror("错误", f"模板不存在:\n{tpl}")
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(tpl, read_only=True)
            ws = wb.active
            hr = int(self.header_row_var.get() or 3)
            cols = []
            for cell in ws[hr]:
                if cell.value:
                    cols.append(str(cell.value).strip())
            wb.close()
            self.msg_var.set(f"模板共 {len(cols)} 列")
            self.preview_text.delete("1.0", "end")
            self.preview_text.insert("end",
                "模板列名（第 {} 行）:\n\n".format(hr))
            for i, c in enumerate(cols, 1):
                self.preview_text.insert("end", f"  {i:>3}. {c}\n")
        except Exception as e:
            messagebox.showerror("读取失败", str(e))

    def _preview(self):
        products = self._get_products(limit=3)
        if not products:
            messagebox.showinfo("提示", "无可抽取的产品")
            return
        try:
            cfg = self._build_config()
            txt = gg.format_preview(products, cfg, max_products=3)
            self.preview_text.delete("1.0", "end")
            self.preview_text.insert("end", txt)
            self.msg_var.set(f"已预览 {len(products)} 组")
        except Exception as e:
            import traceback
            messagebox.showerror("预览失败",
                f"{e}\n\n{traceback.format_exc()[:500]}")

    def _generate(self):
        products = self._get_products()
        if not products:
            messagebox.showinfo("提示",
                "无可抽取的产品\n\n"
                "可能原因:\n"
                "· 数据库为空\n"
                "· 随机抽取模式下剩余未用为 0（点【🗑 清空标记】重置）")
            return
        cfg = self._build_config()

        if not Path(cfg["template_path"]).exists():
            messagebox.showerror("错误",
                f"模板不存在:\n{cfg['template_path']}")
            return
        if not Path(cfg["keyword_file"]).exists():
            if not messagebox.askyesno(
                "关键词文件缺失",
                f"未找到:\n{cfg['keyword_file']}\n\n继续则不加关键词。是否继续?"):
                return

        total = len(products)
        n_files = (total + cfg["max_products_per_file"] - 1) // cfg["max_products_per_file"]
        mode = self.db_sample_mode_var.get()
        mode_txt = f"（随机抽取 {total:,} 条）" if mode == "随机抽取" else "（全部）"
        mark_txt = ""
        if mode == "随机抽取" and self.mark_used_var.get():
            mark_txt = "\n生成后将自动标记为已用（下次抽取自动排除）"
        price_txt = (f"利润=成本×{int(cfg['profit_rate']*100)}%, "
                     f"折扣={cfg['discount']}, "
                     f"费率={cfg['platform_fee']}, "
                     f"运费{cfg['freight_min']}~{cfg['freight_max']}")
        if not messagebox.askyesno(
            "确认生成",
            f"将处理 {total:,} 组产品 {mode_txt}\n"
            f"价格公式: {price_txt}\n"
            f"预计生成 {n_files} 份 xlsx\n"
            f"输出目录: {cfg['output_dir']}"
            f"{mark_txt}\n\n继续?"):
            return

        self.msg_var.set("生成中...")
        self.app.root.update_idletasks()

        # 记录本次生成的产品（用于标记）
        products_snapshot = list(products)
        mark_after = (mode == "随机抽取"
                      and self.mark_used_var.get())

        def worker():
            try:
                paths, stats = gg.generate_xlsx(products_snapshot, cfg)
                self.app.root.after(
                    0, lambda: self._after_generate(
                        paths, stats, products_snapshot, mark_after))
            except Exception as e:
                import traceback
                tb = traceback.format_exc()
                self.app.root.after(0, lambda: messagebox.showerror(
                    "生成失败", f"{e}\n\n{tb[:800]}"))
                self.app.root.after(0,
                    lambda: self.msg_var.set("生成失败"))

        threading.Thread(target=worker, daemon=True).start()

    def _after_generate(self, paths, stats, products, mark_after):
        marked = 0
        if mark_after:
            try:
                om_list = [p["onlymulu"] for p in products
                           if p.get("onlymulu")]
                marked = self._mark_used(om_list)
            except Exception as e:
                print(f"[G] 标记失败: {e}")
        stats["marked"] = marked

        self._on_done(paths, stats)
        self._update_sample_info()

    def _on_done(self, paths, stats):
        lines = [
            "✅ 生成完成",
            f"产品组数: {stats['products']:,}",
            f"文件数:   {stats['files']}",
            f"总行数:   {stats['total_rows']:,}",
        ]
        if stats.get("marked"):
            lines.append(f"新增标记: {stats['marked']:,}")
            lines.append(f"剩余未用: "
                         f"{self._total_products() - self._used_count():,}")
        if stats.get("warnings"):
            lines.append("")
            lines.append(f"⚠️ 首图警告 {len(stats['warnings'])} 条（前 5）:")
            for w in stats["warnings"][:5]:
                lines.append(f"  · {w}")
        if stats.get("missing_cols"):
            lines.append("")
            lines.append("⚠️ 以下列名在模板中未找到（已跳过）:")
            for c in stats["missing_cols"]:
                lines.append(f"  · {c}")
        if stats.get("errors"):
            lines.append("")
            lines.append(f"⚠️ 表达式错误 {len(stats['errors'])} 条（前 5）:")
            for e in stats["errors"][:5]:
                lines.append(f"  · {e}")

        msg = f"✅ 生成 {stats['files']} 份 / {stats['total_rows']} 行"
        if stats.get("marked"):
            msg += f" / 标记 {stats['marked']}"
        self.msg_var.set(msg)
        messagebox.showinfo("完成", "\n".join(lines))

        try:
            out = Path(self._build_config()["output_dir"])
            if out.exists():
                os.startfile(str(out))
        except Exception:
            pass

    def _clear(self):
        self.preview_text.delete("1.0", "end")
        self.msg_var.set("")

    # ==========================================================
    def _execute_one(self, task):
        return "skip", "G 模块无单任务执行"