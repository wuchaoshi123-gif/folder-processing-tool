# tabs/tab_f/f_xlsx.py
"""F 菜单 — xlsx 导入导出"""
import shutil
import tkinter as tk
from pathlib import Path
from datetime import datetime
from tkinter import ttk, messagebox, filedialog

try:
    from openpyxl import Workbook, load_workbook
    _XLSX_OK = True
except ImportError:
    _XLSX_OK = False


class FXlsxMixin:
    """xlsx 导入导出功能"""

    def _export_full_xlsx(self):
        if not _XLSX_OK:
            messagebox.showerror("缺少依赖",
                                 "请先安装 openpyxl:\n\npip install openpyxl")
            return
        if not self._conn:
            self._refresh_db()
            if not self._conn:
                return

        p = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx"), ("All", "*.*")],
            initialfile=f"products_full_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            title="导出完整数据库到 xlsx")
        if not p:
            return

        try:
            cur = self._conn.cursor()
            cur.execute("SELECT * FROM products ORDER BY id")
            rows = [dict(r) for r in cur.fetchall()]

            self.set_status(f"正在导出 {len(rows)} 条到 xlsx...")
            count = self._write_xlsx(p, rows)

            size_kb = Path(p).stat().st_size / 1024
            messagebox.showinfo(
                "导出成功",
                f"完整数据库已导出:\n{p}\n\n"
                f"记录数: {count:,}\n"
                f"大小: {size_kb:.1f} KB")
            self.set_status(f"✅ 已导出 {count} 条到 {Path(p).name}")
            print(f"[F] ✅ 完整导出: {p} ({count} 条)")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))
            print(f"[F] ❌ 导出失败: {e}")

    def _export_filtered_xlsx(self):
        if not _XLSX_OK:
            messagebox.showerror("缺少依赖",
                                 "请先安装 openpyxl:\n\npip install openpyxl")
            return
        if not self._all_rows:
            messagebox.showinfo("提示", "请先查询出结果，再导出")
            return

        p = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx"), ("All", "*.*")],
            initialfile=f"products_filtered_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            title="导出当前筛选结果到 xlsx")
        if not p:
            return

        try:
            self.set_status(f"正在导出 {len(self._all_rows)} 条到 xlsx...")
            count = self._write_xlsx(p, self._all_rows)
            size_kb = Path(p).stat().st_size / 1024
            messagebox.showinfo(
                "导出成功",
                f"筛选结果已导出:\n{p}\n\n"
                f"记录数: {count:,}\n"
                f"大小: {size_kb:.1f} KB")
            self.set_status(f"✅ 已导出 {count} 条到 {Path(p).name}")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _write_xlsx(self, path, rows):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("products")
        ws.append(self.XLSX_HEADERS)
        for row in rows:
            values = []
            for col in self.XLSX_COLUMNS:
                v = row.get(col, "")
                if v is None:
                    v = ""
                v = str(v)
                if len(v) > 32000:
                    v = v[:32000]
                values.append(v)
            ws.append(values)
        wb.save(path)
        return len(rows)

    def _import_xlsx(self):
        if not _XLSX_OK:
            messagebox.showerror("缺少依赖",
                                 "请先安装 openpyxl:\n\npip install openpyxl")
            return
        if not self._conn:
            self._refresh_db()
            if not self._conn:
                return

        p = filedialog.askopenfilename(
            title="选择 xlsx 备份文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("All", "*.*")])
        if not p:
            return

        xlsx_path = Path(p)
        if not xlsx_path.exists():
            messagebox.showerror("错误", "文件不存在")
            return

        try:
            self.set_status("正在读取 xlsx...")
            rows = self._read_xlsx(xlsx_path)
        except Exception as e:
            messagebox.showerror("读取失败", f"无法读取 xlsx:\n{e}")
            return

        if not rows:
            messagebox.showinfo("提示", "xlsx 文件为空或格式错误")
            return

        self.set_status(f"正在分析 {len(rows)} 条记录的重复情况...")
        cur = self._conn.cursor()
        cur.execute("SELECT onlymulu FROM products")
        existing = set(r[0] for r in cur.fetchall())

        new_rows, dup_rows, invalid_rows = [], [], []
        for r in rows:
            om = r.get("onlymulu", "").strip()
            if not om:
                invalid_rows.append(r)
                continue
            (dup_rows if om in existing else new_rows).append(r)

        print(f"\n[F] 📥 导入分析:")
        print(f"   xlsx 总记录: {len(rows)}")
        print(f"   新记录: {len(new_rows)}")
        print(f"   重复记录: {len(dup_rows)}")
        print(f"   无效记录: {len(invalid_rows)}")

        action = self._ask_import_action(
            xlsx_path, len(rows),
            len(new_rows), len(dup_rows), len(invalid_rows))

        if action is None:
            print(f"[F] 用户取消导入")
            return

        self._do_import_xlsx(action, new_rows, dup_rows)

    def _read_xlsx(self, path):
        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        iterator = ws.iter_rows(values_only=True)

        try:
            headers = next(iterator)
        except StopIteration:
            wb.close()
            return []

        header_map = {}
        for i, h in enumerate(headers):
            if h is None:
                continue
            h_str = str(h).strip()
            for col_en, col_cn in zip(self.XLSX_COLUMNS, self.XLSX_HEADERS):
                if h_str == col_cn or h_str == col_en:
                    header_map[i] = col_en
                    break

        if not header_map:
            for i, col_en in enumerate(self.XLSX_COLUMNS):
                header_map[i] = col_en

        for row in iterator:
            if row is None:
                continue
            if all(v is None or (isinstance(v, str) and not v.strip())
                   for v in row):
                continue
            record = {}
            for i, v in header_map.items():
                if i < len(row):
                    val = row[i]
                    if val is None:
                        val = ""
                    elif isinstance(val, datetime):
                        val = val.strftime("%Y-%m-%d %H:%M:%S")
                    else:
                        val = str(val).strip() if not isinstance(val, str) else val.strip()
                    record[v] = val
                else:
                    record[v] = ""
            rows.append(record)

        wb.close()
        return rows

    def _ask_import_action(self, xlsx_path, total, new_cnt, dup_cnt, invalid_cnt):
        win = tk.Toplevel(self.app.root)
        win.title("导入 xlsx")
        win.geometry("620x560")
        win.transient(self.app.root)
        win.grab_set()
        result = {"action": None}

        # 标题
        header = ttk.Frame(win, padding=12)
        header.pack(fill="x")
        ttk.Label(header, text="📥 导入 xlsx 数据",
                  font=("", 13, "bold")).pack(anchor="w")
        ttk.Label(header, text=str(xlsx_path),
                  foreground="#666",
                  font=("Consolas", 8)).pack(anchor="w", pady=(4, 0))

        # 分析
        stat = ttk.LabelFrame(win, text=" 文件分析 ", padding=12)
        stat.pack(fill="x", padx=12, pady=6)
        ttk.Label(stat, text=f"xlsx 记录总数: {total:,}",
                  font=("", 10)).pack(anchor="w")
        ttk.Label(stat, text=f"新记录（不存在）: {new_cnt:,}",
                  foreground="#008000",
                  font=("", 10, "bold")).pack(anchor="w", pady=(4, 0))
        ttk.Label(stat, text=f"重复记录（已存在）: {dup_cnt:,}",
                  foreground="#C00000",
                  font=("", 10, "bold")).pack(anchor="w", pady=(4, 0))
        if invalid_cnt > 0:
            ttk.Label(stat, text=f"无效记录（缺唯一目录）: {invalid_cnt:,}",
                      foreground="#CC6600",
                      font=("", 10)).pack(anchor="w", pady=(4, 0))

        # 选项
        choice = ttk.LabelFrame(win, text=" 请选择处理方式 ", padding=12)
        choice.pack(fill="both", expand=True, padx=12, pady=6)
        ttk.Label(choice, text="重复记录（已存在的 onlymulu）如何处理：",
                  font=("", 10, "bold")).pack(anchor="w", pady=(0, 8))

        # 跳过
        opt1 = ttk.Frame(choice)
        opt1.pack(fill="x", pady=6)
        tk.Button(opt1, text="⏭  跳过重复并导入",
                  font=("", 10, "bold"), width=24, height=1,
                  bg="#E8F4FF", activebackground="#C8E4FF",
                  relief="raised", bd=2,
                  command=lambda: self._confirm_action(win, result, "skip"),
                  ).pack(side="left", padx=(0, 10))
        ttk.Label(opt1,
                  text=f"保留现有 {dup_cnt:,} 条数据\n"
                       f"只导入新的 {new_cnt:,} 条记录",
                  foreground="#004080", font=("", 9), justify="left"
                  ).pack(side="left", anchor="w")

        # 覆盖
        opt2 = ttk.Frame(choice)
        opt2.pack(fill="x", pady=6)
        tk.Button(opt2, text="🔄  覆盖重复并导入",
                  font=("", 10, "bold"), width=24, height=1,
                  bg="#FFF7E0", activebackground="#FFE7A0",
                  relief="raised", bd=2,
                  command=lambda: self._confirm_action(win, result, "overwrite"),
                  ).pack(side="left", padx=(0, 10))
        ttk.Label(opt2,
                  text=f"用 xlsx 数据更新\n"
                       f"现有 {dup_cnt:,} 条重复记录",
                  foreground="#804000", font=("", 9), justify="left"
                  ).pack(side="left", anchor="w")

        # 清空后导入
        opt3 = ttk.Frame(choice)
        opt3.pack(fill="x", pady=6)
        tk.Button(opt3, text="🧹  清空后导入",
                  font=("", 10, "bold"), width=24, height=1,
                  bg="#FFE0E0", activebackground="#FFC0C0",
                  relief="raised", bd=2,
                  command=lambda: self._confirm_action(win, result, "wipe"),
                  ).pack(side="left", padx=(0, 10))
        ttk.Label(opt3,
                  text="⚠️ 危险！删除当前库所有数据\n"
                       "完全用 xlsx 替换（不可恢复）",
                  foreground="#C00000", font=("", 9, "bold"), justify="left"
                  ).pack(side="left", anchor="w")

        # 底部
        bottom = ttk.Frame(win, padding=12)
        bottom.pack(fill="x", side="bottom")
        ttk.Label(bottom,
                  text="💡 导入前会自动备份当前数据库",
                  foreground="#666", font=("", 8)).pack(side="left")
        ttk.Button(bottom, text="取消", width=12,
                   command=win.destroy).pack(side="right", padx=4)

        win.wait_window()
        return result["action"]

    def _confirm_action(self, win, result, action):
        if action == "overwrite":
            if not messagebox.askyesno(
                "确认覆盖",
                "将用 xlsx 数据覆盖当前数据库中所有重复记录。\n\n"
                "此操作不可撤销（但会自动备份）。\n\n确定继续?"):
                return
        elif action == "wipe":
            confirm_win = tk.Toplevel(win)
            confirm_win.title("危险操作确认")
            confirm_win.geometry("420x220")
            confirm_win.transient(win)
            confirm_win.grab_set()

            ttk.Label(confirm_win, text="⚠️ 危险操作",
                      foreground="#C00000",
                      font=("", 14, "bold")).pack(pady=(20, 10))
            ttk.Label(confirm_win,
                      text="此操作会清空当前数据库的所有记录！\n\n"
                           "请输入 确认 两字继续：",
                      justify="center").pack(pady=10)

            entry_var = tk.StringVar()
            entry = ttk.Entry(confirm_win, textvariable=entry_var,
                              width=20, justify="center")
            entry.pack(pady=10)
            entry.focus_set()

            def on_ok():
                if entry_var.get().strip() == "确认":
                    confirm_win.destroy()
                    result["action"] = action
                    win.destroy()
                else:
                    messagebox.showwarning("提示", "请输入正确的确认文字")

            btn_frame = ttk.Frame(confirm_win)
            btn_frame.pack(pady=10)
            ttk.Button(btn_frame, text="取消",
                       command=confirm_win.destroy).pack(side="left", padx=6)
            ttk.Button(btn_frame, text="✔ 确认清空",
                       command=on_ok).pack(side="left", padx=6)

            confirm_win.wait_window()
            return

        result["action"] = action
        win.destroy()

    def _do_import_xlsx(self, action, new_rows, dup_rows):
        db_path = Path(self.db_path_var.get().strip())
        t0 = datetime.now()

        backup_path = db_path.parent / (
            f"{db_path.stem}_preimport_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
        try:
            if self._conn:
                self._conn.commit()
            shutil.copy2(db_path, backup_path)
            print(f"[F] 当前数据库已备份: {backup_path.name}")
        except Exception as e:
            print(f"[F] ⚠️ 备份失败: {e}")
            if not messagebox.askyesno("备份失败",
                                       f"无法备份当前数据库:\n{e}\n\n仍要继续导入?"):
                return

        try:
            cur = self._conn.cursor()
            cur.execute("BEGIN")
            inserted = updated = skipped = 0
            errors = []

            def insert_row(row):
                cur.execute("""
                    INSERT INTO products
                        (nums, onlymulu, englishname, local_url, web_url, rufanyi,
                         shoessize, chuprice, images, bands,
                         bandsizechart, source_dir, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    row.get("nums") or None,
                    row.get("onlymulu"),
                    row.get("englishname", ""),
                    row.get("local_url", ""),
                    row.get("web_url", ""),
                    row.get("rufanyi", ""),
                    row.get("shoessize", ""),
                    row.get("chuprice") or None,
                    row.get("images") or None,
                    row.get("bands", ""),
                    row.get("bandsizechart", ""),
                    row.get("source_dir", ""),
                    row.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                ))

            if action == "wipe":
                cur.execute("DELETE FROM products")
                for row in new_rows + dup_rows:
                    try:
                        insert_row(row)
                        inserted += 1
                    except Exception as e:
                        errors.append(str(e)[:80])
            else:
                for row in new_rows:
                    try:
                        insert_row(row)
                        inserted += 1
                    except Exception as e:
                        errors.append(str(e)[:80])

                if action == "skip":
                    skipped = len(dup_rows)
                elif action == "overwrite":
                    for row in dup_rows:
                        try:
                            cur.execute("""
                                UPDATE products SET
                                    nums = ?, englishname = ?, local_url = ?,
                                    web_url = ?, rufanyi = ?, shoessize = ?,
                                    chuprice = ?, images = ?, bands = ?,
                                    bandsizechart = ?, source_dir = ?,
                                    created_at = ?
                                WHERE onlymulu = ?
                            """, (
                                row.get("nums") or None,
                                row.get("englishname", ""),
                                row.get("local_url", ""),
                                row.get("web_url", ""),
                                row.get("rufanyi", ""),
                                row.get("shoessize", ""),
                                row.get("chuprice") or None,
                                row.get("images") or None,
                                row.get("bands", ""),
                                row.get("bandsizechart", ""),
                                row.get("source_dir", ""),
                                row.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                row.get("onlymulu"),
                            ))
                            updated += 1
                        except Exception as e:
                            errors.append(str(e)[:80])

            self._conn.commit()
            elapsed = (datetime.now() - t0).total_seconds()

            msg = [f"导入完成,耗时 {elapsed:.2f}s", "─" * 40,
                   f"新增记录: {inserted:,}", f"覆盖更新: {updated:,}",
                   f"跳过重复: {skipped:,}"]
            if errors:
                msg.append(f"失败记录: {len(errors):,}")
                msg.append("错误样本（前 5 条）:")
                for e in errors[:5]:
                    msg.append(f"  · {e}")
            msg.append("─" * 40)
            msg.append(f"备份: {backup_path.name}")

            messagebox.showinfo("导入完成", "\n".join(msg))
            self.set_status(
                f"✅ 导入完成: 新增 {inserted} / 覆盖 {updated} / 跳过 {skipped}")
            self._refresh_db()
        except Exception as e:
            try:
                self._conn.rollback()
            except Exception:
                pass
            messagebox.showerror(
                "导入失败",
                f"{e}\n\n当前数据库已回滚\n备份: {backup_path}")