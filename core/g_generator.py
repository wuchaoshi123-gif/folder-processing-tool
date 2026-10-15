# core/g_generator.py
"""
G 模块 —— xlsx 产品表格生成器（v4）
★ v4 新增:
  - build_name: 先确保 rufanyi 含 обувь / кроссовки（缺则前置），再走关键词拼接
"""
import random
import re
from pathlib import Path
from datetime import datetime

try:
    from openpyxl import load_workbook
    _XLSX_OK = True
except ImportError:
    _XLSX_OK = False


# ============================================================
#  默认配置
# ============================================================
DEFAULT_G_CONFIG = {
    "template_path": "平板鞋.xlsx",
    "output_dir": "",
    "header_row": 3,
    "data_start_row": 5,
    "max_products_per_file": 100,

    "img_suffix": "?x-oss-process=style/fnew06",
    "first_img": "1.jpg",
    "sample_count": 0,
    "img_mode": "single",

    "keyword_file": "关键词.txt",
    "keyword_pos": "random",
    "name_max_len": 60,

    # 价格公式参数
    "profit_rate": 1.10,
    "discount": 0.75,
    "platform_fee": 0.17,
    "freight_min": 66,
    "freight_max": 99,

    "fields_a": [
        ["相片", "PHOTO"],
        ["名称", "NAME"],
        ["价格", "calc_price(CHUPRICE)"],
        ["原价", "价格*3"],
        ["组成", "brand_or_empty(BANDS)"],
    ],
    "fields_b": [
        ["团体", "GROUP"],
        ["卖家的文章", "团体"],
        ["尺寸", "SIZE"],
        ["俄罗斯尺寸", "SIZE_INT-1"],
    ],
    "fields_c": [
        ["物品", "板鞋"],
        ["说明书", "Для защиты обуви Наша упаковка удалила коробку для обуви"],
        ["颜色", "Multi"],
        ["每包装数量", "1"],
    ],
}


# ============================================================
#  表达式求值
# ============================================================
def _make_eval_context():
    return {
        "int": int, "float": float, "str": str, "len": len,
        "abs": abs, "round": round, "min": min, "max": max,
        "random": random,
        "rand": lambda a, b: random.randint(int(a), int(b)),
        "rand11": lambda: (datetime.now().strftime("%y%m%d")
                            + "".join(random.choices("0123456789", k=11))),
        "rand_date_prefix": lambda: datetime.now().strftime("%y%m%d"),
        "choice": lambda seq: random.choice(list(seq)),
        "join": lambda sep, items: str(sep).join(str(x) for x in items),
        "brand_or_empty": lambda b: ("" if (not b or str(b).strip().lower() == "allall")
                                     else b),
    }


def _eval_with_scope(expr, scope, on_error=None):
    if expr is None:
        return ""
    expr = str(expr).strip()
    if not expr:
        return ""
    try:
        ctx = _make_eval_context()
        ctx.update(scope)
        return eval(expr, {"__builtins__": {}}, ctx)
    except Exception as e:
        if on_error:
            on_error(f"表达式 {expr!r} 求值失败: {e}")
        return ""


# ============================================================
#  单项构建
# ============================================================
def _is_number(s):
    try:
        float(s)
        return True
    except (ValueError, TypeError):
        return False


def split_sizes(shoessize):
    if not shoessize:
        return []
    parts = re.split(r"[,，、\s]+", str(shoessize).strip())
    return [p.strip() for p in parts if p.strip()]


def _fname_of(url):
    """从 URL 提取文件名（去目录、去参数、小写）"""
    return url.split("?")[0].split("/")[-1].lower()


def build_photo_string(web_url, bandsizechart, config, warnings=None):
    """
    ★ v6 极简版：
      1. 解析所有 URL（自动补后缀）
      2. 找出首图（默认 1.jpg）放第一
      3. 找不到 1.jpg → 用编号最小的图当首图
      4. 其他图随机打乱
      5. 末尾追加尺码表
    """
    suffix = (config.get("img_suffix") or "").strip()
    first_name = (config.get("first_img") or "1.jpg").strip().lower()
    sample = int(config.get("sample_count") or 0)

    # ---- 1. 解析所有 URL ----
    urls = [u.strip() for u in (web_url or "").split("\n") if u.strip()]
    if suffix:
        urls = [(u if "?" in u else u + suffix) for u in urls]

    # ---- 2. 分离首图 / 其他图 ----
    first_url = None
    others = []
    for u in urls:
        if _fname_of(u) == first_name and first_url is None:
            first_url = u
        else:
            others.append(u)

    # ---- 3. fallback: 没有指定首图 → 用编号最小 ----
    if first_url is None and others:
        def _num(u):
            s = _fname_of(u).rsplit(".", 1)[0]
            return int(s) if s.isdigit() else 999999
        others.sort(key=_num)
        first_url = others.pop(0)
        if warnings is not None:
            warnings.append(
                f"未找到 {first_name}，改用 {_fname_of(first_url)} 作首图")

    # ---- 4. 其他图随机 ----
    if 0 < sample < len(others):
        others = random.sample(others, sample)
    random.shuffle(others)

    # ---- 5. 拼接: 首图 + 其他 + 尺码表 ----
    parts = []
    if first_url:
        parts.append(first_url)
    parts.extend(others)
    if bandsizechart and bandsizechart.strip():
        parts.append(bandsizechart.strip())

    return ";".join(parts)


# ============================================================
#  ★ 名称预处理：补全 обувь / кроссовки
# ============================================================
# 这两个词是俄语电商核心搜索词（обувь=鞋类, кроссовки=运动鞋）
# 若 rufanyi 里缺少，就在最前面补上
KEY_TERMS = ("обувь", "кроссовки")


def _ensure_key_terms(rufanyi):
    """
    ★ 检查 rufanyi 是否包含 обувь 和 кроссовки
      缺哪个就在前面补哪个（保持前后顺序 обувь кроссовки）
    """
    if not rufanyi:
        return rufanyi
    lower = rufanyi.lower()
    prefix = []
    for term in KEY_TERMS:
        if term not in lower:
            prefix.append(term)
    if prefix:
        return " ".join(prefix) + " " + rufanyi
    return rufanyi


def _truncate_by_word(text, max_len):
    """按单词边界截断到不超过 max_len"""
    if not text or len(text) <= max_len:
        return text
    words = text.split()
    result = []
    cur_len = 0
    for w in words:
        add = len(w) + (1 if result else 0)
        if cur_len + add > max_len:
            break
        result.append(w)
        cur_len += add
    if not result:
        return text[:max_len]
    return " ".join(result)


def build_name(rufanyi, keyword, position, max_len):
    """
    ★ v4 流程:
      1. 先在 rufanyi 前面补 обувь / кроссовки（缺哪个补哪个）
      2. 再走原来的逻辑：随机关键词 + rufanyi 拼接
      3. 最后按 max_len 单词边界截断
    """
    rufanyi = (rufanyi or "").strip()

    # ★ 第 1 步：补全关键搜索词
    rufanyi = _ensure_key_terms(rufanyi)

    keyword = (keyword or "").strip()

    if not keyword:
        name = rufanyi
        if max_len and len(name) > max_len:
            name = _truncate_by_word(name, max_len)
        return " ".join(name.split())

    # 给"关键词 + 1 个空格"留位置
    kw_space = len(keyword) + 1
    if max_len and max_len <= kw_space:
        return keyword[:max_len]

    ru_max = max_len - kw_space if max_len else None

    if position == "before":
        use_before = True
    elif position == "after":
        use_before = False
    else:
        use_before = random.random() < 0.5

    if ru_max and len(rufanyi) > ru_max:
        rufanyi = _truncate_by_word(rufanyi, ru_max)

    if use_before:
        name = f"{keyword} {rufanyi}".strip()
    else:
        name = f"{rufanyi} {keyword}".strip()

    name = " ".join(name.split())
    if max_len and len(name) > max_len:
        name = _truncate_by_word(name, max_len)
    return name


# ============================================================
#  加载关键词
# ============================================================
def load_keywords(path):
    p = Path(path)
    if not p.exists():
        return []
    result = []
    with open(p, "r", encoding="utf-8-sig", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#"):
                result.append(line)
    return result


# ============================================================
#  ★ 价格公式
# ============================================================
def make_price_calculator(profit_rate, discount, platform_fee,
                          freight_min, freight_max):
    denom = discount * (1 - platform_fee)
    if denom <= 0:
        denom = 1e-9

    def calc_price(cost, freight=None):
        try:
            c = float(cost)
        except (TypeError, ValueError):
            return 0
        if freight is None:
            f = random.randint(int(freight_min), int(freight_max))
        else:
            try:
                f = float(freight)
            except (TypeError, ValueError):
                f = 0
        return round((c * (1 + float(profit_rate)) + f) / denom)

    return calc_price


# ============================================================
#  生成一个产品的多行
# ============================================================
def generate_product_rows(p, config, keywords, errors=None):
    on_error = (errors.append if errors is not None else None)

    rufanyi = p.get("rufanyi", "") or ""
    englishname = p.get("englishname", "") or ""
    chuprice = p.get("chuprice") or 0
    web_url = p.get("web_url", "") or ""
    bandsizechart = p.get("bandsizechart", "") or ""
    shoessize = p.get("shoessize", "") or ""
    bands = p.get("bands", "") or ""

    profit_rate = float(config.get("profit_rate", 1.10))
    discount = float(config.get("discount", 0.75))
    platform_fee = float(config.get("platform_fee", 0.17))
    freight_min = int(config.get("freight_min", 66))
    freight_max = int(config.get("freight_max", 99))

    calc_price = make_price_calculator(
        profit_rate, discount, platform_fee, freight_min, freight_max)

    photo = build_photo_string(web_url, bandsizechart, config)
    keyword = random.choice(keywords) if keywords else ""
    name_default = build_name(
        rufanyi, keyword,
        config.get("keyword_pos", "random"),
        int(config.get("name_max_len") or 60),
    )
    group_id = (datetime.now().strftime("%y%m%d")
                + "".join(random.choices("0123456789", k=11)))

    a_scope = {
        "RUFANYI": rufanyi,          # ★ 保持原始值（不带补词）
        "ENGLISHNAME": englishname,
        "CHUPRICE": chuprice,
        "WEBURL": web_url,
        "BANDSIZECHART": bandsizechart,
        "SHOESSIZE": shoessize,
        "BANDS": bands,
        "PHOTO": photo,
        "KEYWORD": keyword,
        "NAME": name_default,        # ★ 已含补词 + 关键词
        "GROUP": group_id,
        "PROFIT_RATE": profit_rate,
        "DISCOUNT": discount,
        "PLATFORM_FEE": platform_fee,
        "FREIGHT_MIN": freight_min,
        "FREIGHT_MAX": freight_max,
        "calc_price": calc_price,
    }

    a_fields = {}
    for col, expr in config.get("fields_a", []):
        scope = {**a_scope, **a_fields}
        a_fields[col] = _eval_with_scope(expr, scope, on_error)

    sizes = split_sizes(shoessize) or [""]
    rows = []
    for size in sizes:
        # ★ SIZE_INT: 整数尺码返回 int, 带小数返回 float (不再截断)
        if _is_number(size):
            f = float(size)
            size_int = int(f) if f == int(f) else f
            size_float = f
        else:
            size_int = None
            size_float = None
        b_scope = {
            **a_scope,
            **a_fields,
            "SIZE": size,
            "SIZE_INT": size_int,  # 36 → 36,  36.5 → 36.5
            "SIZE_FLOAT": size_float,  # 永远返回 float
        }

        b_fields = {}
        for col, expr in config.get("fields_b", []):
            scope = {**b_scope, **b_fields}
            b_fields[col] = _eval_with_scope(expr, scope, on_error)

        row = {}
        row.update(a_fields)
        row.update(b_fields)
        for col, val in config.get("fields_c", []):
            row[col] = val
        rows.append(row)

    return rows


# ============================================================
#  xlsx 读写
# ============================================================
def read_header_map(ws, header_row):
    col_map = {}
    for col_idx, cell in enumerate(ws[header_row], 1):
        val = cell.value
        if val is None:
            continue
        name = str(val).strip()
        if name and name not in col_map:
            col_map[name] = col_idx
    return col_map


def write_rows_to_sheet(ws, col_map, rows, start_row):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

    missing_cols = set()
    for i, row_data in enumerate(rows):
        excel_row = start_row + i
        for col_name, value in row_data.items():
            col_idx = col_map.get(col_name)
            if not col_idx:
                missing_cols.add(col_name)
                continue
            try:
                ws.cell(row=excel_row, column=col_idx, value=value)
            except Exception:
                pass
    return missing_cols


# ============================================================
#  完整流程
# ============================================================
def generate_xlsx(products, config, progress_cb=None):
    if not _XLSX_OK:
        raise RuntimeError("缺少 openpyxl，请执行: pip install openpyxl")

    template_path = Path(config["template_path"])
    if not template_path.exists():
        raise FileNotFoundError(f"模板不存在: {template_path}")

    output_dir = Path(config.get("output_dir") or template_path.parent)
    output_dir.mkdir(parents=True, exist_ok=True)

    keyword_file = config.get("keyword_file") or "关键词.txt"
    keywords = load_keywords(keyword_file)
    if progress_cb:
        progress_cb(f"已加载 {len(keywords)} 个关键词")

    header_row = int(config.get("header_row") or 3)
    start_row = int(config.get("data_start_row") or 5)
    max_per_file = int(config.get("max_products_per_file") or 100)

    errors = []
    all_product_rows = []
    for i, p in enumerate(products):
        rows = generate_product_rows(p, config, keywords, errors)
        all_product_rows.append(rows)
        if progress_cb and (i + 1) % 20 == 0:
            progress_cb(f"生成中 {i + 1}/{len(products)}...")

    out_paths = []
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    all_missing = set()

    for chunk_idx in range(0, len(all_product_rows), max_per_file):
        chunk = all_product_rows[chunk_idx:chunk_idx + max_per_file]
        rows = []
        for product_rows in chunk:
            rows.extend(product_rows)

        wb = load_workbook(str(template_path))
        ws = wb.active
        col_map = read_header_map(ws, header_row)
        missing = write_rows_to_sheet(ws, col_map, rows, start_row)
        all_missing.update(missing)

        file_no = chunk_idx // max_per_file + 1
        out_path = output_dir / f"G_产品表格_{ts}_{file_no:03d}.xlsx"
        wb.save(str(out_path))
        out_paths.append(out_path)

        if progress_cb:
            progress_cb(
                f"第 {file_no} 份已保存: {out_path.name} ({len(chunk)} 组产品)")

    stats = {
        "products": len(products),
        "files": len(out_paths),
        "total_rows": sum(len(pr) for pr in all_product_rows),
        "errors": errors[:20],
        "missing_cols": sorted(all_missing),
    }
    return out_paths, stats


# ============================================================
#  预览
# ============================================================
def format_preview(products, config, max_products=3, max_value_len=100):
    keywords = load_keywords(config.get("keyword_file") or "关键词.txt")
    lines = []
    for i, p in enumerate(products[:max_products]):
        rows = generate_product_rows(p, config, keywords)
        lines.append(f"═══ 产品 {i + 1}: {p.get('onlymulu', '')} ═══")
        lines.append(f"    原 shoessize: {p.get('shoessize', '')}")
        lines.append(f"    共 {len(rows)} 行")
        for j, row in enumerate(rows):
            lines.append(f"  ── 行 {j + 1} ──")
            for col, val in row.items():
                v = str(val)
                if len(v) > max_value_len:
                    v = v[:max_value_len] + f"... (+{len(v) - max_value_len})"
                lines.append(f"    {col:<20} = {v}")
        lines.append("")
    return "\n".join(lines)