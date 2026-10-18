# core/g1_generator.py
import random
import re
import math
import os
from pathlib import Path
from datetime import datetime
import concurrent.futures as cf

try:
    from openpyxl import load_workbook

    _XLSX_OK = True
except ImportError:
    _XLSX_OK = False

# ============================================================
#  默认配置 (G1 专属)
# ============================================================
DEFAULT_G1_CONFIG = {
    "template_dir": "",
    "output_dir": "",
    "header_row": 1,
    "data_start_row": 2,
    "max_products_per_file": 100,

    "img_suffix": "",
    "first_img": "1.jpg",

    "keyword_file": "关键词.txt",
    "keyword_pos": "random",
    "name_max_len_ru": 255,
    "name_max_len_en": 128,

    "profit_rate": 1.10,
    "discount": 0.75,
    "platform_fee": 0.17,
    "freight_min": 66,
    "freight_max": 99,

    # 自定义描述前缀
    "desc_html_prefix": """<h1><strong>Заметка</strong></h1>
<h3><em>Наша посылка будет отправлена в почтовое отделение в России через почтовую
системуПосылка прибудет в ваше местное российское почтовое
отделениеПожалуйста свяжитесь с местным почтовым отделением чтобы получить
вашу посылку.</em></h3>
<h3><em>Пожалуйста не выбирайте отдельно &quot;ящик для обуви&quot;.Это не
правильный размер обуви и наша фабрика не может отправить его отдельноЕсли
вам нужна пара обуви с коробкойВы можете выбрать &quot;ящик для обуви&quot;
+ правильный размерЭто правильный способ совершить покупку</em></h3>""",

    "fields_a": [
        ["Model (seller SPU ID)*", "GROUP"],
        ["Product name (Russian)", "NAME_RU"],
        ["Product name (English)*", "NAME_EN"],
        ["Main image*", "MAIN_IMAGE"],
        ["Image #2", "IMAGE_2"],
        ["Image #3", "IMAGE_3"],
        ["Image #4", "IMAGE_4"],
        ["Image #5", "IMAGE_5"],
        ["Image #6", "IMAGE_6"],
        ["Description (Russian)", "DESC_HTML"],  # 已恢复合并后的 DESC_HTML 变量
        ["Discounted price, CNY", "calc_price(CHUPRICE)"],
        ["Price, CNY*", "Discounted_price_CNY*3"],
        ["组成", "brand_or_empty(BANDS)"],
    ],
    "fields_b": [
        ["SKU ID", "SKU_ID"],
        ["Size", "SIZE"],
    ],
    "fields_c": [
        ["物品", "板鞋"],
        ["说明书", "Для защиты обуви Наша упаковка удалила коробку для обуви"],
        ["颜色", "Multi"],
        ["每包装数量", "1"],
    ],
}


# ============================================================
#  表达式求值
# ============================================================
def _make_eval_context():
    return {
        "int": int, "float": float, "str": str, "len": len,
        "abs": abs, "round": round, "min": min, "max": max,
        "random": random,
        "rand": lambda a, b: random.randint(int(a), int(b)),
        "rand11": lambda: (datetime.now().strftime("%y%m%d")
                           + "".join(random.choices("0123456789", k=11))),
        "rand_date_prefix": lambda: datetime.now().strftime("%y%m%d"),
        "choice": lambda seq: random.choice(list(seq)),
        "join": lambda sep, items: str(sep).join(str(x) for x in items),
        "brand_or_empty": lambda b: ("" if (not b or str(b).strip().lower() == "allall")
                                     else b),
    }


def _eval_with_scope(expr, scope, on_error=None):
    if expr is None: return ""
    expr = str(expr).strip()
    if not expr: return ""
    try:
        ctx = _make_eval_context()
        ctx.update(scope)
        return eval(expr, {"__builtins__": {}}, ctx)
    except Exception as e:
        if on_error: on_error(f"表达式 {expr!r} 求值失败: {e}")
        return ""


def _add_safe_aliases(scope):
    for k, v in list(scope.items()):
        safe_k = str(k).replace(' ', '_').replace(',', '').replace('，', '')
        if safe_k != k:
            scope[safe_k] = v
    return scope


# ============================================================
#  工具函数
# ============================================================
def _is_number(s):
    try:
        float(s); return True
    except:
        return False


def split_sizes(shoessize):
    if not shoessize: return []
    parts = re.split(r"[,，、\s]+", str(shoessize).strip())
    return [p.strip() for p in parts if p.strip()]


def _truncate_by_word(text, max_len):
    if not text or len(text) <= max_len: return text
    words = text.split()
    result, cur_len = [], 0
    for w in words:
        add = len(w) + (1 if result else 0)
        if cur_len + add > max_len: break
        result.append(w);
        cur_len += add
    return " ".join(result) if result else text[:max_len]


KEY_TERMS = ("обувь", "кроссовки")


def _ensure_key_terms(rufanyi):
    if not rufanyi: return rufanyi
    lower = rufanyi.lower()
    prefix = [term for term in KEY_TERMS if term not in lower]
    return " ".join(prefix) + " " + rufanyi if prefix else rufanyi


def build_names(rufanyi, englishname, keyword, position, max_len_ru, max_len_en):
    ru = _ensure_key_terms((rufanyi or "").strip())
    kw = (keyword or "").strip()
    if kw:
        kw_space = len(kw) + 1
        ru_max = max_len_ru - kw_space if max_len_ru > kw_space else max_len_ru
        use_before = random.random() < 0.5 if position == "random" else (position == "before")
        if len(ru) > ru_max: ru = _truncate_by_word(ru, ru_max)
        ru_name = f"{kw} {ru}".strip() if use_before else f"{ru} {kw}".strip()
    else:
        ru_name = ru
    if max_len_ru and len(ru_name) > max_len_ru: ru_name = _truncate_by_word(ru_name, max_len_ru)

    en_name = (englishname or "").strip()
    if max_len_en and len(en_name) > max_len_en: en_name = _truncate_by_word(en_name, max_len_en)
    return " ".join(ru_name.split()), " ".join(en_name.split())


def load_keywords(path):
    p = Path(path)
    if not p.exists(): return []
    return [line.strip() for line in open(p, "r", encoding="utf-8-sig", errors="ignore") if
            line.strip() and not line.startswith("#")]


def make_price_calculator(profit_rate, discount, platform_fee, freight_min, freight_max):
    denom = discount * (1 - platform_fee)
    if denom <= 0: denom = 1e-9

    def calc_price(cost, freight=None):
        try:
            c = float(cost)
        except:
            return 0
        f = random.randint(int(freight_min), int(freight_max)) if freight is None else float(freight)
        return round((c * (1 + float(profit_rate)) + f) / denom)

    return calc_price


# ============================================================
#  速卖通图片与 HTML 拼接
# ============================================================
def process_ae_images(web_url, bandsizechart, config):
    suffix = (config.get("img_suffix") or "").strip()
    first_name = (config.get("first_img") or "1.jpg").strip().lower()

    urls = [u.strip() for u in (web_url or "").split("\n") if u.strip()]
    if suffix: urls = [(u if "?" in u else u + suffix) for u in urls]

    size_chart = (bandsizechart or "").strip()
    if size_chart and suffix:
        size_chart = size_chart if "?" in size_chart else size_chart + suffix

    main_img, others = None, []
    for u in urls:
        fname = u.split("?")[0].split("/")[-1].lower()
        if fname == first_name and main_img is None:
            main_img = u
        else:
            others.append(u)

    if main_img is None and others:
        def _num(u):
            s = u.split("?")[0].split("/")[-1].rsplit(".", 1)[0]
            return int(s) if s.isdigit() else 999999

        others.sort(key=_num)
        main_img = others.pop(0)

    random.shuffle(others)
    img2 = others[0] if len(others) > 0 else ""
    img3 = others[1] if len(others) > 1 else ""
    img4 = others[2] if len(others) > 2 else ""
    img5 = others[3] if len(others) > 3 else ""
    img6 = size_chart

    # 拼接图片部分的 HTML
    all_imgs_for_html = list(urls)
    random.shuffle(all_imgs_for_html)
    if size_chart: all_imgs_for_html.append(size_chart)

    images_html = ""
    for u in all_imgs_for_html:
        images_html += f'\n<p><img src="{u}"></p>'

    # 将自定义文字和图片 HTML 拼接
    html_prefix = config.get("desc_html_prefix", "")
    desc_html = html_prefix + images_html

    return main_img or "", img2, img3, img4, img5, img6, desc_html, images_html


# ============================================================
#  生成产品行
# ============================================================
def generate_product_rows(p, config, keywords, errors=None):
    on_error = (errors.append if errors is not None else None)

    rufanyi = p.get("rufanyi", "") or ""
    englishname = p.get("englishname", "") or ""
    chuprice = p.get("chuprice") or 0
    web_url = p.get("web_url", "") or ""
    bandsizechart = p.get("bandsizechart", "") or ""
    shoessize = p.get("shoessize", "") or ""
    bands = p.get("bands", "") or ""

    calc_price = make_price_calculator(
        float(config.get("profit_rate", 1.10)), float(config.get("discount", 0.75)),
        float(config.get("platform_fee", 0.17)), int(config.get("freight_min", 66)),
        int(config.get("freight_max", 99))
    )

    main_img, img2, img3, img4, img5, img6, desc_html, images_html = process_ae_images(web_url, bandsizechart, config)

    keyword = random.choice(keywords) if keywords else ""
    name_ru, name_en = build_names(rufanyi, englishname, keyword, config.get("keyword_pos", "random"),
                                   int(config.get("name_max_len_ru") or 255), int(config.get("name_max_len_en") or 128))

    group_id = (datetime.now().strftime("%y%m%d") + "".join(random.choices("0123456789", k=11)))

    a_scope = {
        "RUFANYI": rufanyi, "ENGLISHNAME": englishname, "CHUPRICE": chuprice,
        "WEBURL": web_url, "BANDSIZECHART": bandsizechart, "SHOESSIZE": shoessize,
        "BANDS": bands, "KEYWORD": keyword, "GROUP": group_id,
        "NAME_RU": name_ru, "NAME_EN": name_en,
        "MAIN_IMAGE": main_img, "IMAGE_2": img2, "IMAGE_3": img3,
        "IMAGE_4": img4, "IMAGE_5": img5, "IMAGE_6": img6,
        "DESC_HTML": desc_html,
        "HTML_PREFIX": config.get("desc_html_prefix", ""),
        "IMAGES_HTML": images_html,
        "calc_price": calc_price,
    }

    a_fields = {}
    for col, expr in config.get("fields_a", []):
        scope = _add_safe_aliases({**a_scope, **a_fields})
        a_fields[col] = _eval_with_scope(expr, scope, on_error)

    sizes = split_sizes(shoessize) or [""]
    rows = []
    for idx, size in enumerate(sizes, start=1):
        size_int = int(float(size)) if _is_number(size) and float(size).is_integer() else (
            float(size) if _is_number(size) else None)
        b_scope = {
            **a_scope, **a_fields, "SIZE": size, "SIZE_INT": size_int,
            "SKU_ID": f"{group_id}-{idx}",
        }

        b_fields = {}
        for col, expr in config.get("fields_b", []):
            scope = _add_safe_aliases({**b_scope, **b_fields})
            b_fields[col] = _eval_with_scope(expr, scope, on_error)

        row = {**a_fields, **b_fields}
        for col, val in config.get("fields_c", []):
            row[col] = val
        rows.append(row)

    return rows


# ============================================================
#  多进程任务与分发
# ============================================================
def _worker_generate_file(args):
    template_path, output_path, products, config, keywords = args
    try:
        errors, all_rows = [], []
        for p in products:
            all_rows.extend(generate_product_rows(p, config, keywords, errors))

        wb = load_workbook(str(template_path))
        ws = wb.active
        header_row, start_row = int(config.get("header_row") or 1), int(config.get("data_start_row") or 2)

        col_map = {str(cell.value).strip(): col_idx for col_idx, cell in enumerate(ws[header_row], 1) if cell.value}
        if ws.max_row >= start_row: ws.delete_rows(start_row, ws.max_row - start_row + 1)

        missing = set()
        for i, row_data in enumerate(all_rows):
            r_idx = start_row + i
            for c_name, val in row_data.items():
                c_idx = col_map.get(c_name)
                if not c_idx:
                    missing.add(c_name);
                    continue
                ws.cell(row=r_idx, column=c_idx, value=val)

        wb.save(str(output_path))
        return True, str(output_path), list(missing), errors[:5]
    except Exception as e:
        return False, str(output_path), [], [str(e)]


def generate_xlsx_multi(products, config, progress_cb=None):
    if not _XLSX_OK: raise RuntimeError("缺少 openpyxl")

    tpl_dir, out_dir = Path(config["template_dir"]), Path(config["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)

    templates = [f for f in tpl_dir.glob("*.xlsx") if not f.name.startswith("~$")]
    if not templates: raise ValueError("未找到任何模板")

    keywords = load_keywords(config.get("keyword_file") or "关键词.txt")
    max_per = int(config.get("max_products_per_file") or 100)

    total_files = math.ceil(len(products) / max_per)
    base_per_tpl = total_files // len(templates)
    remainder = total_files % len(templates)

    tasks, p_idx = [], 0
    for i, tpl in enumerate(templates):
        files_this_tpl = base_per_tpl + (1 if i < remainder else 0)
        for f_idx in range(1, files_this_tpl + 1):
            chunk = products[p_idx: p_idx + max_per]
            p_idx += len(chunk)
            tasks.append((tpl, out_dir / f"{tpl.stem}_{f_idx}.xlsx", chunk, config, keywords))

    if progress_cb: progress_cb(f"开始并发生成 {len(tasks)} 个文件...")

    success, all_missing, all_errors = 0, set(), []
    with cf.ProcessPoolExecutor(max_workers=min(os.cpu_count() or 4, 8)) as ex:
        futures = {ex.submit(_worker_generate_file, t): t for t in tasks}
        for idx, fut in enumerate(cf.as_completed(futures)):
            ok, path, miss, errs = fut.result()
            if ok:
                success += 1
            else:
                all_errors.extend(errs)
            all_missing.update(miss)
            if progress_cb: progress_cb(f"已生成 {idx + 1}/{len(tasks)} 份")

    return success, len(tasks), list(all_missing), all_errors


def format_preview(products, config, max_products=3, max_value_len=100):
    keywords = load_keywords(config.get("keyword_file") or "关键词.txt")
    lines = []
    for i, p in enumerate(products[:max_products]):
        rows = generate_product_rows(p, config, keywords)
        lines.append(f"═══ 产品 {i + 1}: {p.get('onlymulu', '')} ═══")
        for j, row in enumerate(rows):
            lines.append(f"  ── 行 {j + 1} ──")
            for col, val in row.items():
                v = str(val)
                if len(v) > max_value_len: v = v[:max_value_len] + f"... (+{len(v) - max_value_len})"
                lines.append(f"    {col:<20} = {v}")
        lines.append("")
    return "\n".join(lines)